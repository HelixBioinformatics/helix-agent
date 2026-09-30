#!/usr/bin/env python3
"""Build the pre-inference expert workbook from archived, technically checked inputs.
No expected labels or human paraphrases are generated. Requires openpyxl.
"""
import csv,hashlib,json,shutil,zipfile
from copy import copy
from pathlib import Path
from datetime import datetime,timezone
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break
ROOT=Path(__file__).resolve().parents[2]
BASE=ROOT/'evaluation/trp/corpus-20260930'
OUT=ROOT/'docs/tesis/evaluacion-experta-20260930/para-asesora'
DARK='214E60';BLUE='EAF3F7';YELLOW='FFF2CC';GRAY='F0F3F5';INK='203544'
existing=OUT/'evaluacion.xlsx'
if existing.exists():
 previous=load_workbook(existing)
 for ws in previous:
  for cells in ws:
   for cell in cells:
    if cell.fill.fgColor.rgb=='00FFF2CC' and cell.value is not None:
     raise SystemExit('Refusing to overwrite a workbook containing expert responses: '+str(existing))
wb=Workbook();wb.remove(wb.active)
input_cells=[];validations=[]
def sheet(name,title,subtitle):
 s=wb.create_sheet(name);s.sheet_view.showGridLines=False;s.sheet_view.zoomScale=85;s.freeze_panes='A3'
 for col,width in [('A',27),('B',65),('C',70)]:s.column_dimensions[col].width=width
 s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.orientation='portrait';s.page_setup.paperSize=s.PAPERSIZE_A4;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
 s.page_margins.left=s.page_margins.right=.25;s.page_margins.top=s.page_margins.bottom=.35
 s.oddFooter.center.text=f'{name} · &P / &N'
 full(s,title,'title',34);full(s,subtitle,'blue',60);return s

def style(c,kind='body'):
 c.font=Font(name='Calibri',size=13 if kind=='title' else 11,color='FFFFFF' if kind=='title' else INK,bold=kind in ('title','label'))
 c.fill=PatternFill('solid',fgColor={'title':DARK,'blue':BLUE,'input':YELLOW,'label':GRAY}.get(kind,'FFFFFF'))
 c.alignment=Alignment(vertical='top',wrap_text=True)
 thin=Side(style='hair',color='D7E0E5');c.border=Border(left=thin,right=thin,top=thin,bottom=thin)

def full(s,text,kind='body',height=52):
 r=s.max_row+1 if s.cell(1,1).value is not None else 1
 s.merge_cells(start_row=r,start_column=1,end_row=r,end_column=3);c=s.cell(r,1,text);style(c,kind);s.row_dimensions[r].height=height;return r

def row(s,label,guide,answer=None,height=57,options=None):
 r=s.max_row+1
 for col,value,kind in [(1,label,'label'),(2,guide,'body'),(3,answer,'input' if answer is None else 'body')]:style(s.cell(r,col,value),kind)
 s.row_dimensions[r].height=height
 if answer is None:input_cells.append((s.title,f'C{r}'))
 if options:
  v=DataValidation(type='list',formula1='"'+','.join(options)+'"',allow_blank=True);v.errorTitle='Seleccione una opción';v.error='Use una opción de la lista; explique cualquier duda en la justificación.';v.showErrorMessage=True;s.add_data_validation(v);v.add(s.cell(r,3));validations.append((s.title,f'C{r}',options))
 return r

def header(s):
 r=s.max_row+1
 for i,v in enumerate(['Campo','Qué se espera / información','Su respuesta'],1):style(s.cell(r,i,v),'label')
 s.row_dimensions[r].height=26

def jump(s,rowno,target,label):
 c=s.cell(rowno,3,label);c.hyperlink=f"#'{target[0]}'!A{target[1]}";c.font=Font(name='Calibri',size=11,color='0563C1',underline='single');c.alignment=Alignment(wrap_text=True,vertical='top')

selected=json.loads((BASE/'selected.json').read_text());assert len(selected)==50
for d in selected:
 if d['task']=='detection':d['files']=[f for f in d['files'] if not f['path'].endswith('/unidades.json')]
negatives=json.loads((BASE/'negative-requests.json').read_text());assert len(negatives)==20
case_records=[]
for i,d in enumerate(selected):
 acc=d['acc'];directory='datos/'+acc
 if d['task']=='detection' and d['kind']=='pdb':
  wording=[
   'Quiero buscar posibles repeticiones en {pdb}, cadena {chain}. Usa STRPsearch sobre {directory}/{pdb}.cif y conserva la numeración de autor. Antes de preparar y ejecutar el análisis, muéstrame las acciones y pide mi autorización. Entrega los intervalos candidatos y explica las limitaciones; no los presentes como unidades ya validadas.',
   'Analiza con STRPsearch la cadena {chain} de la estructura {pdb} que está en {directory}/{pdb}.cif. Necesito un listado de regiones y unidades repetitivas candidatas, referido a la numeración de autor. Consulta conmigo antes de las etapas que ejecuten herramientas y conserva la procedencia del archivo.',
   'A partir del archivo {directory}/{pdb}.cif, propón una búsqueda de repeticiones estructurales para la cadena {chain} de {pdb}. Utiliza STRPsearch, solicita mi aprobación antes de preparar la entrada y antes de detectar, e informa los límites predichos con su numeración y sus limitaciones.',
   '¿Puedes buscar unidades repetitivas candidatas en {pdb}, cadena {chain}, con STRPsearch? El archivo es {directory}/{pdb}.cif. Quiero aprobar las etapas antes de que las ejecutes y recibir los intervalos en numeración de autor junto con la información necesaria para identificar la estructura utilizada.',
   'Necesito detectar candidatos a repetición en la cadena {chain} de {pdb}, disponible en {directory}/{pdb}.cif. Prepara el análisis mediante STRPsearch, con aprobación previa a cada etapa de ejecución. Informa regiones, unidades candidatas y limitaciones; la ausencia de candidatos también debe quedar registrada.'
  ][i%5].format(**d,directory=directory)
 elif d['task']=='detection':
  wording=(f"Busca regiones y unidades repetitivas candidatas con STRPsearch en el modelo {d['entry']}, cadena {d['chain']}, contenido en {directory}/modelo.cif. "
   f"Comprueba su identidad con {directory}/procedencia-afdb.json. Antes de preparar la entrada y antes de ejecutar la detección, muéstrame la propuesta y solicita mi aprobación. "
   "Conserva la numeración del modelo y entrega los intervalos candidatos, la procedencia y las limitaciones. No interpretes los valores pLDDT como factores B experimentales ni presentes la detección como una validación física o biológica.")
 else:
  bounds=', '.join(f"{u['start']}-{u['end']}" for u in d['units'])
  wording=(f"Para el modelo {d['entry']}, cadena {d['chain']}, calcula con GeomeTRe la geometría de los tres intervalos suministrados {bounds}. "
   f"Usa {directory}/modelo.cif, la procedencia de {directory}/procedencia-afdb.json, la matriz {directory}/pae.json y los límites de {directory}/unidades.json. "
   "Antes de ejecutar, comprueba la correspondencia de secuencia y numeración. Exijo pLDDT de al menos 70 en todos los residuos seleccionados y PAE de como máximo 5 Å en ambos sentidos entre cada par de unidades. "
   "Muéstrame el plan y pide mi aprobación. Si alguna condición no se cumple, detente y explica la causa. Si se cumplen, entrega los descriptores y su procedencia. Los intervalos provienen de anotaciones de secuencia; no afirmes que este análisis valida biológicamente sus límites ni la estabilidad física del plegamiento.")
 case_records.append({'key':acc,'kind':d['kind'],'task':d['task'],'request_draft':wording,'intended_stratum':'feasible_conditional_plan','input':d,'paraphrases':[None,None],'human_reference':None})
for n,text in enumerate(negatives,1):case_records.append({'key':f'negative-{n:02d}','kind':'scope','request_draft':text,'intended_stratum':'out_of_scope_or_invalid_request','input':None,'paraphrases':[None,None],'human_reference':None})
case_records.sort(key=lambda r:hashlib.sha256(('helix-expert-display-20260930|'+r['key']).encode()).hexdigest())
for n,c in enumerate(case_records,1):c['case_id']=f'S{n:03d}'

s=sheet('Inicio','HELIX · Evaluación experta de solicitudes','Primera entrega de la evaluación real: solicitudes y referencia humana, antes de ejecutar el asistente. Complete las celdas amarillas. Los archivos de apoyo están en datos/, junto a este Excel.')
full(s,'Propósito','title',27)
full(s,'Necesitamos establecer qué debería entender y hacer Helix ante solicitudes de análisis de proteínas. Su criterio servirá como referencia para evaluar después sus respuestas. Aquí se le pide revisar solicitudes, formular dos maneras alternativas de expresar las tareas factibles y describir un plan o respuesta aceptable.',height=72)
full(s,'Este archivo contiene 70 solicitudes propuestas: 50 orientadas a tareas disponibles y 20 para evaluar límites o peticiones inválidas. Están mezcladas y la decisión de cada una le corresponde a usted. Los textos iniciales fueron preparados con asistencia de IA; todavía no son una referencia validada por una persona. Los 50 casos con datos corresponden a proteínas humanas; valore si esa cobertura es adecuada.',height=65)
full(s,'Cómo trabajar','title',27)
full(s,'1. Complete su identificación y lea «Guía» y «Alcance».\n2. En «Índice», abra cada solicitud y su ficha de datos.\n3. En «Solicitudes», revise la redacción, escriba la referencia y reformule las tareas factibles dos veces con sus propias palabras.\n4. Anote dudas y casos pendientes en «Comentarios». Puede devolver avances parciales.',height=100)
full(s,'Se puede trabajar solo con este Excel y la carpeta datos/. No necesita instalar Helix, programar ni ejecutar análisis. Las fichas describen los insumos; los archivos completos permiten consultar dudas. No modifique los datos originales: indique cualquier corrección necesaria en su respuesta.',height=63)
full(s,'Esta entrega aún no contiene respuestas de Helix ni conversaciones para clasificar. Se las entregaremos después de cerrar las referencias y realizar las sesiones. No se le pide simular conversaciones ni fabricar intervenciones humanas.',height=59)
header(s)
for a,b,h in [('Nombre o seudónimo','Identificador con el que desea que se registren sus juicios.',35),('Fecha de entrega','Día/mes/año. Puede actualizarla al devolver el archivo.',35),('Experiencia pertinente','Breve descripción de su experiencia en bioinformática o estructuras de proteínas.',55),('Relación con el proyecto','Indique su papel de asesora y qué partes del informe revisó. Confirme o corrija que no participó en el desarrollo del software.',62),('Uso de sus anotaciones','Indique si acepta que sus respuestas se utilicen para calcular y reportar resultados en la tesis. Puede escribir condiciones o dudas.',58)]:row(s,a,b,height=h)
full(s,'Guarde y devuelva sus primeras decisiones antes de discutirlas con el tesista. La revisión posterior de desacuerdos se conserva por separado. Deje pendientes los casos que no pueda juzgar y explique por qué.',height=54)

s=sheet('Guía','GUÍA · Cómo construir la referencia','Evalúe el significado de cada solicitud y lo que constituiría una respuesta aceptable. No se espera que calcule los valores de los descriptores ni que prediga qué devolverá Helix.')
full(s,'Revisión del texto y paráfrasis','title',27)
full(s,'Para cada solicitud, confirme si la redacción es natural y comprensible. Si necesita cambios, escriba la versión corregida en «Solicitud definitiva»; de lo contrario, escriba «Mantener propuesta». Su referencia debe corresponder a esa versión definitiva.',height=58)
full(s,'Para cada tarea que considere factible, redacte dos paráfrasis propias: otras maneras naturales de pedir exactamente el mismo objetivo, con las mismas entradas, parámetros, condiciones y aprobación. No copie mecánicamente el original. No suprima condiciones para abreviar. Para peticiones fuera de alcance puede escribir «No aplica» en ambas paráfrasis.',height=73)
full(s,'Los campos amarillos se dejan vacíos para recoger su criterio. No hay respuestas esperadas ocultas ni etiquetas automáticas. Si la tarea necesita una aclaración, indique la pregunta concreta y la información que permitiría resolverla antes de usar el caso.',height=58)
full(s,'Factibilidad','title',27)
for a,b in [('Factible','La solicitud permite preparar un plan dentro del alcance. Una verificación de confianza puede dar resultado negativo: detenerse y explicar la causa puede ser la respuesta correcta de un plan condicional.'),('Requiere aclaración','La tarea podría ser viable, pero falta una decisión del usuario que cambia las entradas, operaciones o parámetros. Escriba la pregunta que haría.'),('Fuera de alcance','El objetivo no está cubierto, o exige una interpretación científicamente inválida que no debe ejecutarse tal como está. Explique el límite sin sustituir silenciosamente el objetivo.'),('No evaluable','No puede juzgar el caso por falta de datos o por requerir otra especialidad. Indique qué necesita; no es un fallo automático del agente.')]:row(s,a,b,'Solo lectura',70)
full(s,'Qué constituye una referencia aceptable','title',27)
full(s,'Debe conservar objetivo, proteína, cadena, numeración, operaciones, parámetros y sus fuentes, condiciones para detenerse, aprobación humana y producto final. Puede admitir varias soluciones. No exija una frase exacta ni considere correcto un resultado solo porque recibió aprobación.',height=64)
full(s,'Ejemplo de redacción, ajeno a estos casos','title',27)
full(s,'Ante «Dime qué cadenas contiene la estructura que adjunté, sin analizar geometría», una referencia podría indicar: consultar el archivo, identificar las cadenas y reportar el listado con su fuente; si el archivo no está disponible, solicitarlo. No calcular geometría ni inventar cadenas. No es necesario escribir comandos.',height=66)

s=sheet('Alcance','ALCANCE · Operaciones y datos','Use estas capacidades para decidir qué puede solicitarse al asistente en esta evaluación. Si considera que alguna premisa no basta para responder, indíquelo en el caso correspondiente.')
full(s,'Operaciones disponibles','title',27)
full(s,'• Recuperar estructuras PDB.\n• Detectar regiones y unidades repetitivas candidatas mediante STRPsearch sobre una cadena de una estructura.\n• Calcular geometría con GeomeTRe sobre intervalos suministrados.\n• Comprobar identidad, numeración y datos de confianza cuando correspondan, preparar un plan y solicitar aprobación antes de ejecutar.',height=105)
full(s,'La detección informa candidatos; no certifica repeticiones reales. La geometría calcula descriptores sobre los intervalos dados; no demuestra que sus límites sean una anotación estructural correcta. No se presupone que la salida de cualquier detector se pueda encadenar automáticamente con otra herramienta sin comprobaciones.',height=65)
full(s,'Límites','title',27)
full(s,'No están disponibles como tareas integradas: GWAS, RNA-seq, ensamblaje genómico, diagnóstico, dosis clínicas, docking, dinámica molecular, química cuántica, entrenamiento o ejecución de AlphaFold como predictor nuevo, ni ReUPred. Los límites de RepeatsDB no se convierten automáticamente sin resolver cadena y numeración.',height=68)
full(s,'Tampoco puede demostrarse estabilidad física, función o corrección de un ensamblaje genómico solo con una estructura, con pLDDT alto o con PAE bajo. Los factores B experimentales no son pLDDT.',height=50)
full(s,'Cómo leer los archivos','title',27)
full(s,'Cada caso con datos tiene una subcarpeta datos/ACCESIÓN/. La ficha en «Datos casos» muestra proteína, organismo, cadena, procedencia y archivos. «uniprot.json» contiene anotaciones públicas de secuencia, que pueden apoyarse en predicción o similitud incluso si la entrada está revisada.',height=62)
full(s,'Para PDB: el archivo .cif contiene coordenadas experimentales y «entidad-rcsb.json» identifica la proteína y cadenas. Para modelos predichos: «modelo.cif» contiene coordenadas, «procedencia-afdb.json» identifica modelo y secuencia, «pae.json» contiene la matriz de confianza y «unidades.json» describe los intervalos suministrados.',height=69)
full(s,'Para las tareas de geometría sobre modelos predichos se suministran tres segmentos tomados de las anotaciones Repeat de UniProt. Son entradas del análisis, no una referencia independiente de repeticiones estructurales. Sus límites se han comprobado contra la numeración de una secuencia canónica completa; si duda de su pertinencia, indíquelo.',height=67)
full(s,'Glosario','title',27)
for a,b in [('Numeración de autor','Identificadores de residuos en la estructura experimental; no deben confundirse con posiciones consecutivas de secuencia.'),('pLDDT','Confianza local del modelo, en escala 0 a 100. El umbral 70 es una condición del ejercicio, no una certificación biológica.'),('PAE','Error alineado predicho en ångströms (Å). Se considera el máximo de ambos bloques dirigidos entre cada par de unidades, con umbral 5 Å.'),('Aprobación','La persona autoriza una acción concreta tras revisar su propuesta. Pedir aprobación no equivale a haberla recibido.')]:row(s,a,b,'Solo lectura',57)

index=sheet('Índice','ÍNDICE · Casos para revisar','Pulse «Abrir solicitud» para ir a una ficha. En cada solicitud encontrará un enlace a sus datos. Los identificadores permanecen estables para poder comentar y devolver avances.')
for i,v in enumerate(['Caso','Entrada / tema','Acceso a la ficha'],1):style(index.cell(3,i,v),'label')
requests=sheet('Solicitudes','SOLICITUDES · Referencia y reformulaciones','Complete una ficha por caso. Las celdas amarillas son su respuesta; la propuesta original permanece visible para conservar qué se revisó. El alcance y las condiciones se describen en «Guía» y «Alcance».')
data=sheet('Datos casos','DATOS · Fichas y archivos de apoyo','La carpeta datos/ debe estar junto al Excel. Las fichas corresponden a archivos públicos descargados para esta evaluación. Puede consultarlos; no debe ejecutarlos ni editar los originales.')

fields=[
('Revisión de redacción','Seleccione si mantendría la solicitud o necesita cambios. «No evaluable» si no puede decidir.',['Mantener','Corregir','No evaluable']),
('Solicitud definitiva','Escriba «Mantener propuesta» o una versión corregida completa. Conserve la intención; si no es válida, explique el problema sin convertirla silenciosamente en otra tarea.',None),
('Objetivo esperado','Describa qué quiere obtener la persona, sin añadir una nueva pregunta científica.',None),
('Entradas e identidad','Indique modelo o estructura, cadena, intervalos, archivos y comprobaciones de identidad necesarios. No complete información faltante por su cuenta.',None),
('Operaciones y orden','Describa las acciones apropiadas y su orden. Si corresponde detenerse o rechazar la petición, indique cómo responder.',None),
('Numeración de residuos','Indique el marco de residuos y las correspondencias necesarias. Puede citar límites y fichas; no debe reconstruir un mapa a mano.',None),
('Confianza y condiciones','Anote qué umbrales y comprobaciones se requieren y qué ocurre si no se cumplen. No presuponga que los datos pasan el filtro.',None),
('Parámetros y origen','Registre parámetros, unidades y procedencia de cada valor: solicitud, ficha o aclaración necesaria. Puede remitir a archivos y límites concretos.',None),
('Aprobación humana','Indique cuándo debe pedirse autorización y qué se debe mostrar antes de cada acción.',None),
('Producto y limitaciones','Describa qué salida sería aceptable, incluyendo una negativa justificada si corresponde, y qué afirmaciones no estarían respaldadas.',None),
('Aclaraciones necesarias','Escriba las preguntas que haría antes de decidir. Si no son necesarias, escriba «Ninguna».',None),
('Factibilidad','Use las definiciones de «Guía». El juicio lo establece usted, aunque la propuesta inicial sugiera una tarea disponible.',['Factible','Requiere aclaración','Fuera de alcance','No evaluable']),
('Justificación','Fundamente el juicio, alternativas admisibles y cualquier duda científica. Para un caso no trabajado, escriba «Pendiente».',None),
('Paráfrasis 1','Para una tarea factible, escriba otra solicitud natural con el mismo significado, entradas y condiciones. Redacción humana. Para fuera de alcance: «No aplica».',None),
('Paráfrasis 2','Escriba una segunda formulación distinta que conserve la intención y todas las condiciones de la solicitud definitiva. Para fuera de alcance: «No aplica».',None),
('Equivalencia de sus textos','Confirme si ambas paráfrasis conservan el significado de la solicitud definitiva. No es una valoración de salidas de Helix.',['Sí','Requiere revisión','No aplica']),
('Tiempo empleado (min)','Minutos aproximados de trabajo activo, incluida la revisión y las reformulaciones.',None)]
cellmap={}
for c in case_records:
 cid=c['case_id'];d=c['input']
 if requests.max_row>2:requests.row_breaks.append(Break(id=requests.max_row))
 start=full(requests,f'Solicitud {cid}','title',27)
 full(requests,c['request_draft'],'blue',175 if d and d['kind']=='afdb' else 105)
 datastart=full(data,f'Datos de {cid}','title',27)
 if d:
  full(data,f"{d['protein']} · {d['organism']}\nUniProt {d['acc']} · {d['length']} residuos en secuencia canónica · cadena {d['chain']}\n"+(f"Estructura experimental PDB {d['pdb']}, modelo 1." if d['kind']=='pdb' else f"Modelo predicho {d['entry']}, modelo 1."),height=90)
  if d['task']=='predicted_geometry':
   full(data,'Intervalos suministrados: '+', '.join(f"{u['start']}-{u['end']} ({u['description']})" for u in d['units'])+'. Extremos inclusivos. Numeración del modelo alineada a la secuencia canónica completa. No se suministran inserciones. Origen: anotaciones Repeat de UniProt, sujetas a su revisión.',height=88)
  else:full(data,'No se suministran límites como entrada del detector: la tarea consiste en buscar candidatos. Se conserva la numeración de la estructura y su identidad de proteína/cadena. La ficha identifica si la fuente es experimental o predicha; la pertinencia científica queda para su revisión.',height=60)
  for f in d['files']:
   r=row(data,Path(f['path']).name,'Archivo de apoyo. '+({'uniprot.json':'Anotación de secuencia y sus evidencias.','entidad-rcsb.json':'Identidad de proteína y cadenas.','unidades.json':'Límites y condiciones para las tareas de geometría; no son salida del detector.','procedencia-afdb.json':'Identidad, versión y secuencia del modelo.','pae.json':'Matriz PAE.'}.get(Path(f['path']).name,'Coordenadas de la estructura.')),f['path'],45)
   data.cell(r,3).hyperlink=f['path'];data.cell(r,3).font=Font(name='Calibri',size=11,color='0563C1',underline='single')
  full(data,'Fuentes públicas: https://www.uniprot.org/uniprotkb/'+d['acc']+' y '+('https://www.rcsb.org/structure/'+d['pdb'].upper() if d['kind']=='pdb' else 'https://alphafold.ebi.ac.uk/entry/'+d['acc']),height=45)
 else:full(data,'Esta solicitud se evalúa con el alcance descrito en «Alcance». No tiene archivos asociados. Si considera que faltan datos, indique cuáles y si resolverlos permitiría atender el objetivo solicitado.',height=66)
 r=row(requests,'Datos de apoyo','Abra la ficha correspondiente y consulte datos/ si lo necesita.','Abrir datos',33);jump(requests,r,('Datos casos',datastart),'Abrir datos de '+cid)
 header(requests);cellmap[cid]={}
 for label,guide,options in fields:
  r=row(requests,label,guide,height=80 if label.startswith('Paráfrasis') or label=='Solicitud definitiva' else 53 if label!='Tiempo empleado (min)' else 30,options=options);cellmap[cid][label]=f'C{r}'
 ir=index.max_row+1
 for col,val in [(1,cid),(2,(d['protein']+' · '+d['acc']) if d else c['request_draft'])]:style(index.cell(ir,col,val),'label' if col==1 else 'body')
 index.row_dimensions[ir].height=51;jump(index,ir,('Solicitudes',start),'Abrir solicitud '+cid)
 c['worksheet']='Solicitudes';c['cells']=cellmap[cid];c['human_review_status']='pending';c['source_status']='technical-input-check-only' if d else 'scope-request'

s=sheet('Comentarios','COMENTARIOS · Revisión de la entrega','Registre aquí dificultades generales. Puede devolver el archivo por partes; los casos sin respuesta permanecerán pendientes.')
header(s)
for label,guide in [('Casos terminados','Liste los identificadores revisados o indique un intervalo, por ejemplo S001 a S010.'),('Casos pendientes','Indique qué queda por revisar y cualquier dificultad que impida hacerlo.'),('Problemas con los datos','Indique archivo y caso, y qué información falta o parece incoherente.'),('Cambios en la guía','Proponga cambios en campos o instrucciones que no permitan expresar su criterio.'),('Realismo y cobertura','¿Las solicitudes representan tareas útiles? Señale formulaciones artificiales y tareas relevantes que falten.'),('Tiempo total (min)','Tiempo aproximado de lectura y trabajo, sin pausas largas.'),('Observaciones adicionales','Cualquier otro comentario para mejorar la evaluación.')]:row(s,label,guide,height=80)
full(s,'La siguiente entrega contendrá las respuestas de Helix y, por separado, turnos de conversaciones reales. No clasifique conversaciones inventadas para sustituirlas. Su referencia inicial se conservará antes de mostrarle los resultados.',height=60)

for s in wb:
 s.print_options.horizontalCentered=True;s.print_area=f'A1:C{s.max_row}'
 for r in s:
  for cell in r:
   if isinstance(cell.value,str):cell.value=cell.value.replace('—',':')
wb.save(OUT/'evaluacion.xlsx')
(BASE/'cases.jsonl').write_text(''.join(json.dumps(c,ensure_ascii=False)+'\n' for c in case_records))
(BASE/'workbook-cell-map.json').write_text(json.dumps(cellmap,ensure_ascii=False,indent=2))
manifest=[]
for d in selected:
 for f in d['files']:
  p=OUT/f['path'];assert hashlib.sha256(p.read_bytes()).hexdigest()==f['sha256'];manifest.append({**f,'accession':d['acc']})
(OUT/'datos/fuentes.json').write_text(json.dumps({'retrieved_at':datetime.now(timezone.utc).isoformat(),'sources':manifest,'licenses':{'PDB':'CC0; https://www.wwpdb.org/about/usage-policies','UniProt':'CC BY 4.0; https://www.uniprot.org/help/license','AlphaFold DB':'CC BY 4.0; https://alphafold.ebi.ac.uk/faq'},'scope':'Inputs for expert review; not independent biological ground truth'},ensure_ascii=False,indent=2))
# Archive only reviewed package content, not unrelated failed retrievals.
archive=OUT.parent/'Helix-evaluacion-experta.zip'
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
 for p in [OUT/'evaluacion.xlsx',OUT/'datos/fuentes.json']+[OUT/f['path'] for f in manifest]:z.write(p,p.relative_to(OUT))
loaded=load_workbook(OUT/'evaluacion.xlsx')
for sn,cell in input_cells:assert loaded[sn][cell].value is None
assert len(case_records)==70 and len({r['input']['acc'] for r in case_records if r['input']})==50
assert all('—' not in str(c.value) for s in loaded for row_ in s for c in row_)
for s in loaded:
 for row_ in s:
  for c in row_:
   if c.hyperlink and not c.hyperlink.target.startswith('#'):assert (OUT/c.hyperlink.target).is_file()
report={'status':'prepared_for_human_review_not_frozen','organism_scope':'Homo sapiens','requests':70,'intended_feasible':50,'task_counts':{'detection':25,'predicted_geometry':25},'scope_requests':20,'distinct_input_accessions':50,'human_paraphrases_completed':0,'human_paraphrases_required_for_feasible_intents':100,'human_reference_labels_completed':0,'agent_inference_calls':0,'human_conversation_turns':0,'blank_response_cells':len(input_cells),'workbook_sha256':hashlib.sha256((OUT/'evaluacion.xlsx').read_bytes()).hexdigest(),'data_files':len(manifest),'package_sha256':hashlib.sha256(archive.read_bytes()).hexdigest()}
(BASE/'validation.json').write_text(json.dumps(report,indent=2));print(json.dumps(report,indent=2))
